# -*- coding: utf-8 -*-
"""
換算表 xlsx -> 主要食材(公克) JSON 腳本
========================================
用途：讀取「食譜食材公克換算表.xlsx」的「換算明細」工作表，
     依「連續列」分組（同一道菜的食材在表格中是相鄰的列），將每項食材
     輸出成「食材名稱 換算公克公克」的字串，格式與原始「清洗後食譜.json」
     的 主要食材 欄位一致。

     注意：不同菜名相同（例如兩道都叫「海綿蛋糕」但食材不同）的食譜，
     只要在表格中不是相鄰列，就會被視為「不同的兩筆」而分開輸出，
     不會被合併成一筆。

用法：
    python xlsx轉json.py

輸出範例：
[
  {
    "菜名": "脆皮雞蛋糕",
    "主要食材": [
      "奶油 50公克",
      "牛奶 50公克",
      ...
    ]
  },
  ...
]
"""

import json
from itertools import groupby
from openpyxl import load_workbook

# ========== 參數設定 ==========
INPUT_XLSX = "食譜食材公克換算表.xlsx"
OUTPUT_JSON = "食譜食材公克換算.json"
SHEET_NAME = "換算明細"


def fmt_grams(g):
    """把公克數格式化成字串：整數不顯示小數點，否則保留到小數第二位。"""
    if g is None:
        return None
    g = float(g)
    if g == int(g):
        return str(int(g))
    return str(round(g, 2))


def main():
    wb = load_workbook(INPUT_XLSX, read_only=True)
    ws = wb[SHEET_NAME]

    rows = [r for r in ws.iter_rows(values_only=True)]
    header = rows[0]
    idx = {name: i for i, name in enumerate(header)}
    data_rows = [r for r in rows[1:] if r[idx["菜名"]] is not None]

    # 依「連續列的菜名」分組，同名但不相鄰的菜視為不同的兩筆，不會被合併
    data = []
    for dish, group in groupby(data_rows, key=lambda r: r[idx["菜名"]]):
        items = []
        for row in group:
            name = row[idx["食材名稱"]]
            grams = row[idx["換算公克"]]
            if not name:
                continue
            if grams is not None:
                text = f"{name} {fmt_grams(grams)}公克"
            else:
                # 沒有換算出公克數的項目，保留原食材名稱，不加公克數
                text = name
            items.append(text)
        # 去除同一道菜裡完全重複的「食材名稱+數值」字串
        seen = set()
        deduped = []
        for item in items:
            if item not in seen:
                seen.add(item)
                deduped.append(item)

        data.append({"菜名": dish, "主要食材": deduped})

    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    print(f"共處理 {len(data)} 道菜，已輸出：{OUTPUT_JSON}")


if __name__ == "__main__":
    main()